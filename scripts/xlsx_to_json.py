#!/usr/bin/env python3
"""Convert the PES field-mapping workbook into data.json for the static viewer.
Reads columns by HEADER NAME, not position, so adding/removing/reordering
columns on any module tab (e.g. Beneficiary Type, Visible Tranche, extra
role columns) does not break the build. Role columns are auto-detected as
everything between 'Chain Type' and 'Kobo form field name' on that tab, so
each module tab can have a different number of roles.

Usage: python xlsx_to_json.py <workbook.xlsx> <out.json>
"""
import sys, json
from openpyxl import load_workbook

MODS = ['Application', 'Contract', 'Monitoring Visit', 'Remediation',
        'Payment', 'Beneficiary', 'Project & Activity']

CORE = {
    'field': 'Field name / Action',
    'section': 'Section',
    'sub': 'Sub section',
    'activity': 'Activity Type',
    'status': 'Status',
    'chain': 'Chain Type',
    'kf': 'Kobo form field name',
    'kq': 'Kobo form Question',
    'hide': 'Hide condition / Notes',
    'cov': 'Coverage',
}
# optional columns, included as extra filter dimensions if present on a tab
EXTRA = ['Beneficiary Type', 'Visible Tranche']


def clean(v):
    return '' if v is None else str(v).strip()


def header_map(ws):
    hdr = [clean(c.value) for c in ws[1]]
    idx = {h: i for i, h in enumerate(hdr) if h}
    return hdr, idx


def sheet_rows(ws, module):
    hdr, hidx = header_map(ws)
    missing = [v for v in CORE.values() if v not in hidx]
    if missing:
        print(f'  WARNING [{module}]: missing expected columns {missing} — skipping tab', file=sys.stderr)
        return []
    idx = {k: hidx[v] for k, v in CORE.items()}
    ci = idx['chain']
    ki = idx['kf']
    if ki <= ci:
        print(f'  WARNING [{module}]: "Kobo form field name" is not after "Chain Type" — cannot auto-detect roles, skipping tab', file=sys.stderr)
        return []
    role_names = [hdr[i] for i in range(ci + 1, ki)]
    extra_idx = {e: hidx[e] for e in EXTRA if e in hidx}

    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not clean(r[idx['field']]):
            continue
        roles = [clean(r[ci + 1 + j]) for j in range(len(role_names))]
        row = {
            'mod': module,
            'f': clean(r[idx['field']]),
            'sec': clean(r[idx['section']]),
            'sub': clean(r[idx['sub']]),
            'act': clean(r[idx['activity']]) or 'Any',
            'st': clean(r[idx['status']]) or 'All',
            'ch': clean(r[idx['chain']]) or 'both',
            'roles': roles,
            'kf': clean(r[idx['kf']]),
            'kq': clean(r[idx['kq']]),
            'hide': clean(r[idx['hide']]),
            'cov': clean(r[idx['cov']]),
        }
        for name, i in extra_idx.items():
            row[name] = clean(r[i])
        out.append(row)
    return out, role_names


def main(src, out):
    wb = load_workbook(src, data_only=True)
    rows = []
    role_sets = {}
    for m in MODS:
        name = m[:31]
        if name not in wb.sheetnames:
            continue
        result = sheet_rows(wb[name], m)
        if not result:
            continue
        r, role_names = result
        rows.extend(r)
        role_sets[m] = role_names

    files = []
    if 'Files' in wb.sheetnames:
        ws = wb['Files']
        _, fidx = header_map(ws)
        for r in ws.iter_rows(min_row=2, values_only=True):
            if not r or not clean(r[fidx.get('File / Document Name', 2)]):
                continue
            files.append({
                'mod': clean(r[fidx.get('Module', 1)]),
                'name': clean(r[fidx.get('File / Document Name', 2)]),
                'act': clean(r[fidx.get('Activity Type', 3)]),
                'req': clean(r[fidx.get('Required', 4)]),
                'multi': clean(r[fidx.get('Can upload multiple', 5)]),
                'key': clean(r[fidx.get('Backend Key', 6)]),
            })

    # union of all role names seen, preserving first-seen order, for the
    # role filter dropdown; each row still only carries its own tab's roles
    all_roles = []
    for names in role_sets.values():
        for n in names:
            if n not in all_roles:
                all_roles.append(n)

    json.dump({
        'roles': all_roles,
        'role_sets': role_sets,   # per-module role name lists, same order as each row's `roles` array
        'modules': MODS,
        'rows': rows,
        'files': files,
    }, open(out, 'w'), ensure_ascii=False)
    print(f'wrote {out}: {len(rows)} rows, {len(files)} files, roles per module: '
          + ', '.join(f'{m}={len(role_sets.get(m, []))}' for m in MODS))


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])